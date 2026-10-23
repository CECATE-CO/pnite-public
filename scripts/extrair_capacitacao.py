"""
PNITE — Extração e consolidação dos dados das Capacitações do Transporte Escolar.

Lê os arquivos brutos de cada capacitação (XCTE) em ``data/Capacitação/``:

- ``XCTE - Convocação e Municípios Convidados.html``: polo, data, local e municípios convidados;
- ``XCTE - Formulario (respostas).xlsx``: inscrições;
- ``XCTE - Lista de presença (respostas).xlsx``: presença no evento;
- ``XCTE - Avaliação (respostas).xlsx``: avaliação de reação dos participantes.

As capacitações 1 a 4 não possuem arquivos brutos; seus dados vêm de
``legado_cte_01a04_powerbi.json`` (extraído do modelo do Power BI ``bi dados.pbix``).

Grava as tabelas consolidadas no schema privado ``capacitacao`` do Supabase (migração
``supabase/migrations/0012_capacitacao.sql``), numa única transação, e atualiza a view
pública ``public.capacitacao_dados`` lida pela interface. Os centróides dos municípios
(distância até o polo) vêm de ``pnite_capacitacao_geo.js``, mantido localmente em ``data/``.

Privacidade: nomes, CPF, e-mails, telefones e datas de nascimento são usados
apenas em memória para deduplicar pessoas e ligar presença à inscrição. Nenhum
dado pessoal é gravado no banco; apenas contagens por município e as respostas
das avaliações (anônimas, identificadas por um ID sequencial).

Uso (na raiz do repositório, com o ``.env`` preenchido):
    pip install openpyxl requests python-dotenv
    python scripts/extrair_capacitacao.py
"""

import difflib
import html as html_lib
import json
import math
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from criar_banco_dados_brutos import sql_insert  # noqa: E402
from supabase_api import LIMITE_CORPO, executar_sql, tamanho_corpo  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
PASTA = RAIZ / "data" / "Capacitação"
LISTA_IBGE = RAIZ / "data" / "Transporte Escolar" / "Lista de Municipios IBGE.csv"
GEO_JS = PASTA / "pnite_capacitacao_geo.js"
LEGADO = PASTA / "legado_cte_01a04_powerbi.json"

# Tipo no banco de cada coluna (as demais são text); ver 0012_capacitacao.sql
TIPO_COLUNA = {
    "cod_ibge": "char(7)", "polo_cod_ibge": "char(7)", "uf": "char(2)", "lat": "numeric", "lon": "numeric",
    "numero": "smallint", "capacitacao": "smallint", "data_inicio": "date", "data_fim": "date",
    "convidado": "smallint", "inscritos_total": "integer", "inscritos_gestores": "integer",
    "inscritos_cacs": "integer", "inscritos_outros": "integer", "presentes_total": "integer",
    "presentes_gestores": "integer", "presentes_cacs": "integer", "presentes_outros": "integer",
    "distancia_polo_km": "numeric", "id_pergunta": "smallint", "ordem": "smallint", "nota": "smallint",
}
# Ordem de inserção respeita as chaves estrangeiras
TABELAS_BANCO = ["municipios", "capacitacoes", "municipio_capacitacao", "aval15_perguntas", "aval15_respondentes",
                 "aval15_respostas", "aval616_perguntas", "aval616_respondentes", "aval616_notas", "aval616_canais"]

UF_POR_CODIGO = {"50": "MS", "51": "MT", "52": "GO", "53": "DF", "15": "PA"}

# UFs de onde vieram os municípios convidados de cada capacitação (texto da convocação).
# Usadas para desambiguar nomes iguais em UFs diferentes (ex.: Mundo Novo GO/MS).
UFS_CAPACITACAO = {
    5: ["GO"], 6: ["MS"], 7: ["MS"], 8: ["MT"], 9: ["MT"], 10: ["MT", "GO"],
    11: ["GO", "PA"], 12: ["GO", "DF"], 13: ["MT"], 14: ["GO", "MS", "MT"],
    15: ["GO", "PA"], 16: ["MT"],
}

MESES = {"janeiro": 1, "fevereiro": 2, "março": 3, "marco": 3, "abril": 4, "maio": 5, "junho": 6,
         "julho": 7, "agosto": 8, "setembro": 9, "outubro": 10, "novembro": 11, "dezembro": 12}

# Grafia canônica das respostas fechadas do formulário das capacitações 1 a 5
RESPOSTAS_FECHADAS = {"otimo": "Ótimo", "bom": "Bom", "regular": "Regular", "ruim": "Ruim",
                      "sim": "Sim", "nao": "Não", "em parte": "Em parte"}

PERGUNTAS_6A16 = ["Inscrição", "Divulgação", "Data da formação", "Horário da formação",
                  "Local da formação", "Duração da formação", "Avaliação geral da formação"]

avisos = []


def aviso(msg):
    avisos.append(msg)


# --------------------------------------------------------------------------
# Normalização de textos e municípios
# --------------------------------------------------------------------------
def sem_acento(s):
    return "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))


def limpa(s):
    if s is None:
        return ""
    s = str(s).replace("\xa0", " ").replace("_x000a_", " ")
    s = re.sub(r"\s+", " ", s).strip()
    # Erros de fórmula das planilhas (ex.: PROCV sem correspondência) equivalem a vazio
    return "" if s.upper() in ("#N/A", "#REF!", "#VALUE!", "NONE") else s


def chave_nome(s):
    s = sem_acento(limpa(s)).lower()
    s = re.sub(r"[’'´`]", "", s)
    s = re.sub(r"[^a-z0-9 ]", " ", s)
    return re.sub(r"\s+", " ", s).strip()


NOMES_UF = {"goias": "GO", "mato grosso do sul": "MS", "mato grosso": "MT", "distrito federal": "DF",
            "para": "PA", "go": "GO", "ms": "MS", "mt": "MT", "df": "DF", "pa": "PA"}


def separa_uf(texto):
    """'Baliza (GO)' -> ('Baliza', 'GO'); 'Santana do Araguaia - PA' -> (..., 'PA')."""
    t = limpa(texto)
    m = re.match(r"^(.*?)\s*\(\s*([A-Za-z]{2})\s*\)?\s*$", t)
    if m:
        return m.group(1), m.group(2).upper()
    m = re.match(r"^(.*?)\s*[-/–]\s*([A-Za-zÀ-ú ]+)$", t)
    if m and chave_nome(m.group(2)) in NOMES_UF:
        return m.group(1), NOMES_UF[chave_nome(m.group(2))]
    return t, None


def carrega_municipios():
    """Municípios de GO, MT, MS, DF e PA a partir da lista oficial do IBGE."""
    muns = {}
    with open(LISTA_IBGE, encoding="utf-8-sig") as f:
        next(f)
        for linha in f:
            p = linha.rstrip("\n").split(";")
            # A lista também traz linhas de distritos (código com 9 dígitos)
            if len(p) < 4 or p[0] not in UF_POR_CODIGO or len(p[2]) != 7:
                continue
            muns[p[2]] = {"cod_ibge": p[2], "nome": p[3], "uf": UF_POR_CODIGO[p[0]]}
    return muns


MUNICIPIOS = carrega_municipios()
INDICE_NOME = defaultdict(list)
for _m in MUNICIPIOS.values():
    INDICE_NOME[chave_nome(_m["nome"])].append(_m)
# Grafias alternativas encontradas nos formulários
ALIASES = {"goias velho": "goias", "cidade de goias": "goias", "brasilia df": "brasilia",
           "douradas": "dourados", "santa izabel": "santa isabel"}


def localiza_municipio(texto, ufs_preferidas=None, convidados=None):
    """Retorna o registro IBGE do município ou None. Usa a UF explícita, depois a
    lista de convidados da capacitação e as UFs da capacitação para desambiguar."""
    if not limpa(texto):
        return None
    nome, uf = separa_uf(texto)
    k = chave_nome(nome)
    k = ALIASES.get(k, k)
    cands = INDICE_NOME.get(k, [])
    if not cands:
        todas = [m for lst in INDICE_NOME.values() for m in lst]
        pool = [m for m in todas if (uf and m["uf"] == uf) or (not uf and ufs_preferidas and m["uf"] in ufs_preferidas)]
        perto = difflib.get_close_matches(k, [chave_nome(m["nome"]) for m in pool], n=1, cutoff=0.88)
        if perto:
            cands = [m for m in pool if chave_nome(m["nome"]) == perto[0]]
    if uf:
        cands = [m for m in cands if m["uf"] == uf] or cands
    if len(cands) > 1 and convidados:
        c2 = [m for m in cands if m["cod_ibge"] in convidados]
        if c2:
            cands = c2
    if len(cands) > 1 and ufs_preferidas:
        c2 = [m for m in cands if m["uf"] in ufs_preferidas]
        if c2:
            cands = c2
    return cands[0] if len(cands) == 1 else None


def segmento(texto):
    t = chave_nome(texto)
    if "cacs" in t or "conselh" in t:
        return "CACS"
    if "gest" in t:
        return "Gestores"
    return "Outros" if t else None


def digitos_cpf(v):
    d = re.sub(r"\D", "", limpa(v))
    return d.zfill(11) if 9 <= len(d) <= 11 else ""


# --------------------------------------------------------------------------
# Leitura de planilhas
# --------------------------------------------------------------------------
def le_aba(caminho, nome_aba=None):
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb[nome_aba] if nome_aba else wb.worksheets[0]
    linhas = [r for r in ws.iter_rows(values_only=True) if any(c not in (None, "") for c in r)]
    wb.close()
    if not linhas:
        return [], []
    cab = [limpa(c) for c in linhas[0]]
    return cab, [list(r) + [None] * (len(cab) - len(r)) for r in linhas[1:]]


def colunas(cab, *prefixos):
    """Índices das colunas cujo cabeçalho começa com algum dos prefixos (ordem do arquivo)."""
    out = []
    for i, c in enumerate(cab):
        cl = c.lower()
        if any(cl.startswith(p.lower()) for p in prefixos):
            out.append(i)
    return out


def primeiro(row, idxs):
    for i in idxs:
        if i < len(row) and limpa(row[i]):
            return limpa(row[i])
    return ""


# --------------------------------------------------------------------------
# Convocações (HTML)
# --------------------------------------------------------------------------
def le_convocacao(n):
    t = (PASTA / f"{n}CTE - Convocação e Municípios Convidados.html").read_text(encoding="utf-8")
    texto = html_lib.unescape(re.sub(r"<[^>]+>", " ", t)).replace("\xa0", " ")
    texto = re.sub(r"\s+", " ", texto)
    itens = [limpa(html_lib.unescape(re.sub(r"<[^>]+>", "", x))) for x in re.findall(r"<li[^>]*>(.*?)</li>", t, re.S)]
    itens = [i for i in itens if i]

    # Data: frase oficial da convocação ("no dia 18 de março", "nos dias 24 e 25 de setembro de 2025")
    intro = texto[texto.find("CONVOCA"):][:900]
    m = re.search(r"nos? dias? (\d{1,2})(?:º)?(?: e (\d{1,2}))? de ([a-zç]+)(?: de (\d{4}))?", intro, re.I)
    if not m:
        raise ValueError(f"{n}CTE: data não encontrada na convocação")
    # Sem ano na frase: usa a primeira data dd/mm/aaaa do documento (o HTML pode quebrar dígitos em spans)
    ano = m.group(4) or re.search(r"\d{2}/\d{2}/(20\d{2})", texto.replace(" ", "")).group(1)
    mes = MESES[sem_acento(m.group(3).lower())] if sem_acento(m.group(3).lower()) in MESES else MESES[m.group(3).lower()]
    d1 = date(int(ano), mes, int(m.group(1)))
    d2 = date(int(ano), mes, int(m.group(2))) if m.group(2) else d1

    loc = re.search(r"Local:\s*(.*?)\s*(?:Endereço|Link)", texto)
    local = limpa(loc.group(1)).rstrip(". ") if loc else None
    return {"data_inicio": d1.isoformat(), "data_fim": d2.isoformat(), "local": local, "convidados_texto": itens}


# --------------------------------------------------------------------------
# Inscrições e presença
# --------------------------------------------------------------------------
def pessoas_formulario(caminho, n, convidados, todas_abas=False, registrar_avisos=True):
    """Lista de pessoas (cpf, chave do nome, município, segmento) de um formulário.
    Considera a aba principal e, se pedido, as demais abas que tenham CPF e município."""
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    nomes_abas = wb.sheetnames if todas_abas else wb.sheetnames[:1]
    wb.close()
    pessoas = []
    for aba in nomes_abas:
        cab, linhas = le_aba(caminho, aba)
        i_cpf = colunas(cab, "CPF")
        i_nome = colunas(cab, "Nome Completo", "Nome")
        i_mun = colunas(cab, "Município que representa", "Cidade")
        i_seg = colunas(cab, "Você faz parte do", "Função")
        if not i_mun and not i_cpf:
            continue
        for r in linhas:
            muni_txt = primeiro(r, i_mun)
            mun = localiza_municipio(muni_txt, UFS_CAPACITACAO.get(n), convidados) if muni_txt else None
            if muni_txt and not mun and registrar_avisos:
                aviso(f"{n}CTE [{caminho.name} / {aba}]: município não identificado: '{muni_txt}'")
            pessoas.append({
                "cpf": digitos_cpf(primeiro(r, i_cpf)),
                "nome": chave_nome(primeiro(r, i_nome)),
                "mun": mun["cod_ibge"] if mun else None,
                "seg": segmento(primeiro(r, i_seg)),
                "aba_principal": aba == nomes_abas[0],
            })
    return pessoas


def deduplica(pessoas):
    vistos, out = set(), []
    for p in pessoas:
        k = p["cpf"] or p["nome"]
        if not k or k in vistos:
            continue
        vistos.add(k)
        out.append(p)
    return out


# Pessoas de todos os formulários (por CPF/nome), para localizar o município de quem
# assinou a presença sem informá-lo e não está na inscrição da própria capacitação.
REFERENCIA_GLOBAL = {}


def carrega_referencia_global():
    for n in range(5, 17):
        for arq in (f"{n}CTE - Formulario (respostas).xlsx", f"{n}CTE - Lista de presença (respostas).xlsx"):
            for p in pessoas_formulario(PASTA / arq, n, None, todas_abas=True, registrar_avisos=False):
                if p["mun"]:
                    for k in (p["cpf"], p["nome"]):
                        if k:
                            REFERENCIA_GLOBAL.setdefault(k, p)


def consolida_capacitacao(n, convidados):
    form = PASTA / f"{n}CTE - Formulario (respostas).xlsx"
    pres = PASTA / f"{n}CTE - Lista de presença (respostas).xlsx"

    inscritos_todos = pessoas_formulario(form, n, convidados, todas_abas=True)
    inscritos = deduplica([p for p in inscritos_todos if p["aba_principal"]])
    por_cpf = {p["cpf"]: p for p in inscritos_todos if p["cpf"] and p["mun"]}
    por_nome = {p["nome"]: p for p in inscritos_todos if p["nome"] and p["mun"]}

    presentes_brutos = pessoas_formulario(pres, n, convidados, todas_abas=True)
    # Abas auxiliares da lista de presença (ex.: "Inscrições", "Certificado") servem só de referência cruzada
    for p in presentes_brutos:
        if not p["aba_principal"]:
            if p["cpf"] and p["mun"]:
                por_cpf.setdefault(p["cpf"], p)
            if p["nome"] and p["mun"]:
                por_nome.setdefault(p["nome"], p)
    presentes = deduplica([p for p in presentes_brutos if p["aba_principal"]])
    sem_mun = 0
    for p in presentes:
        ref = (por_cpf.get(p["cpf"]) or por_nome.get(p["nome"])
               or REFERENCIA_GLOBAL.get(p["cpf"]) or REFERENCIA_GLOBAL.get(p["nome"]))
        if not p["mun"] and ref:
            p["mun"] = ref["mun"]
        if not p["seg"] and ref:
            p["seg"] = ref["seg"]
        if not p["seg"]:
            p["seg"] = "Outros"
        if not p["mun"]:
            sem_mun += 1
    if sem_mun:
        aviso(f"{n}CTE: {sem_mun} presente(s) sem município identificável (contados apenas no total da capacitação)")

    agg = defaultdict(lambda: {"inscritos": Counter(), "presentes": Counter()})
    for p in inscritos:
        if p["mun"]:
            agg[p["mun"]]["inscritos"][p["seg"] or "Outros"] += 1
    for p in presentes:
        if p["mun"]:
            agg[p["mun"]]["presentes"][p["seg"]] += 1
    return agg, len(inscritos), len(presentes), sem_mun


# --------------------------------------------------------------------------
# Avaliações
# --------------------------------------------------------------------------
def avaliacao_5():
    """5ª capacitação: mesmo formulário das capacitações 1 a 4 (escala Ótimo/Bom/Regular/Ruim e Sim/Em parte/Não)."""
    cab, linhas = le_aba(PASTA / "5CTE - Avaliação (respostas).xlsx")
    i_pub = colunas(cab, "Você faz parte do")[0]
    out = []
    for seq, r in enumerate(linhas, 1):
        pub = segmento(r[i_pub]) or "Outros"
        idr = f"CAP05-{'GES' if pub == 'Gestores' else 'CAC' if pub == 'CACS' else 'OUT'}-{seq:03d}"
        vistos = set()
        for i, h in enumerate(cab):
            v = limpa(r[i])
            if not v:
                continue
            m = re.match(r"(.*?)\[(.*)\]\s*$", h)
            hl = h.lower()
            if hl.startswith("aspectos gerais") and m:
                item = ("Aspectos Gerais", "Capacitação", limpa(m.group(2)))
            elif hl.startswith("em relação aos resultados") and m:
                sub = limpa(re.sub(r"^Em relação aos resultados do:\s*", "", m.group(1), flags=re.I))
                sub = re.sub(r"\s*:\s+", ": ", sub)
                item = ("Módulos", sub, limpa(m.group(2)))
            elif hl.startswith("sistema eletrônico de gestão") and m:
                item = ("SETE", "Sistema Eletrônico de Gestão do Transporte Escolar", limpa(m.group(2)))
            elif hl.startswith("quanto ao curso ser ministrado") and m:
                item = ("EAD", "Ensino a Distância", limpa(m.group(2)))
            elif hl.startswith("sugestão de temas"):
                item = ("Sugestões", "Temas futuros", "Sugestão de temas para futuras capacitações (Webinar, vídeos de Youtube, postagens de Instagram, etc.)")
            elif hl.startswith("como ficou sabendo"):
                for canal in [c.strip() for c in v.split(",") if c.strip()]:
                    out.append([idr, 5, pub, "Divulgação", "Origem da informação", "Como ficou sabendo da Capacitação?", canal])
                continue
            elif hl.startswith("comentários"):
                item = ("Comentários", "Comentários gerais", "Comentários e/ou Sugestões:")
            else:
                continue  # palestrantes, carimbo de data/hora e público
            if item in vistos:
                continue
            vistos.add(item)
            if item[0] in ("Aspectos Gerais", "Módulos", "SETE", "EAD"):
                v = RESPOSTAS_FECHADAS.get(chave_nome(v), v)
            out.append([idr, 5, pub, *item, v])
    return out


def avaliacao_6a16():
    respostas, notas, canais = [], [], []
    for n in range(6, 17):
        cab, linhas = le_aba(PASTA / f"{n}CTE - Avaliação (respostas).xlsx")
        i_mun = colunas(cab, "Município que representa")[0]
        i_car = colunas(cab, "Você faz parte do")[0]
        i_notas = colunas(cab, "Avaliar de 1 a 5")
        i_pos = colunas(cab, "Indique dois aspectos que você mais gostou")[0]
        i_mel = colunas(cab, "Indique dois aspectos que poderiam")[0]
        i_inst = colunas(cab, "Instituição vinculada")[0]
        i_inst2 = colunas(cab, "Mencionar a")[0]
        i_tem = colunas(cab, "Sugestão de temas")[0]
        i_como = colunas(cab, "Como ficou sabendo")[0]
        i_com = colunas(cab, "Comentários")[0]
        if len(i_notas) != 7:
            raise ValueError(f"{n}CTE: esperadas 7 notas, encontradas {len(i_notas)}")
        for seq, r in enumerate(linhas, 1):
            idr = f"CTE{n:02d}-{seq:04d}"
            muni_txt = limpa(r[i_mun])
            mun = localiza_municipio(muni_txt, UFS_CAPACITACAO.get(n))
            if muni_txt and not mun:
                aviso(f"{n}CTE [Avaliação]: município não identificado: '{muni_txt}'")
            inst = limpa(r[i_inst]).rstrip(";").strip() or None
            if inst and limpa(r[i_inst2]):
                inst = f"{inst} — {limpa(r[i_inst2])}"
            respostas.append([idr, n, mun["cod_ibge"] if mun else None,
                              mun["nome"] if mun else (muni_txt.title() or None),
                              segmento(r[i_car]) or "Outros", inst,
                              limpa(r[i_pos]) or None, limpa(r[i_mel]) or None,
                              limpa(r[i_tem]) or None, limpa(r[i_com]) or None])
            for ordem, i in enumerate(i_notas, 1):
                try:
                    nota = int(float(r[i]))
                except (TypeError, ValueError):
                    continue
                if 1 <= nota <= 5:
                    notas.append([idr, ordem, nota])
            for canal in [c.strip() for c in limpa(r[i_como]).split(",") if c.strip()]:
                canais.append([idr, canal])
    return respostas, notas, canais


# --------------------------------------------------------------------------
# Geometria: centróide dos municípios (para distância ao polo)
# --------------------------------------------------------------------------
def centroides():
    t = GEO_JS.read_text(encoding="utf-8")
    geo = json.loads(t[t.index("{"):t.rindex("}") + 1])
    out = {}
    for f in geo["features"]:
        g = f["geometry"]
        polys = [g["coordinates"]] if g["type"] == "Polygon" else g["coordinates"]
        ax = ay = aa = 0.0
        for poly in polys:
            ring = poly[0]
            for (x0, y0), (x1, y1) in zip(ring, ring[1:]):
                c = x0 * y1 - x1 * y0
                aa += c
                ax += (x0 + x1) * c
                ay += (y0 + y1) * c
        if aa:
            out[f["properties"]["cod"]] = (round(ay / (3 * aa), 4), round(ax / (3 * aa), 4))
    return out


def distancia_km(a, b):
    if not a or not b:
        return None
    (la1, lo1), (la2, lo2) = a, b
    p = math.pi / 180
    h = math.sin((la2 - la1) * p / 2) ** 2 + math.cos(la1 * p) * math.cos(la2 * p) * math.sin((lo2 - lo1) * p / 2) ** 2
    return round(12742 * math.asin(math.sqrt(h)), 1)


# --------------------------------------------------------------------------
# Montagem
# --------------------------------------------------------------------------
def tabela(colunas_, linhas):
    return {"colunas": colunas_, "linhas": linhas}


def gravar_banco(dados):
    """Substitui o conteúdo do schema ``capacitacao`` numa única transação e atualiza a view pública."""
    tabelas = ", ".join(f"capacitacao.{t}" for t in ["metadados", *TABELAS_BANCO])
    partes = ["BEGIN;", f"TRUNCATE {tabelas};"]
    for nome in TABELAS_BANCO:
        t = dados[nome]
        colunas_ = [("seq", "integer")] + [(c, TIPO_COLUNA.get(c, "text")) for c in t["colunas"]]
        partes.append(sql_insert(f"capacitacao.{nome}", colunas_, [[i + 1, *linha] for i, linha in enumerate(t["linhas"])]))
    partes.append(sql_insert("capacitacao.metadados", [("versao_esquema", "smallint"), ("gerado_em", "text"),
                                                       ("fontes", "jsonb"), ("presentes_sem_municipio", "jsonb")],
                             [[dados["versao_esquema"], dados["gerado_em"], json.dumps(dados["fontes"], ensure_ascii=False),
                               json.dumps(dados["presentes_sem_municipio"])]]))
    partes += ["REFRESH MATERIALIZED VIEW public.capacitacao_dados;", "COMMIT;"]
    sql = "\n".join(partes)
    if tamanho_corpo(sql) > LIMITE_CORPO:
        sys.exit(f"Carga da capacitação ({tamanho_corpo(sql):,} bytes) excede o limite de uma requisição ({LIMITE_CORPO:,}).")
    executar_sql(sql)
    contagem = executar_sql("SELECT " + ", ".join(f"(SELECT count(*) FROM capacitacao.{t}) AS {t}" for t in TABELAS_BANCO) + ";")[0]
    divergentes = [t for t in TABELAS_BANCO if contagem[t] != len(dados[t]["linhas"])]
    print(f"Gravado no banco: schema capacitacao ({sum(contagem.values()):,} linhas) e public.capacitacao_dados"
          + (f" — DIVERGENTE em {', '.join(divergentes)}" if divergentes else " — OK"))


def main():
    legado = json.loads(LEGADO.read_text(encoding="utf-8"))
    cent = centroides()

    capacitacoes = []
    fato = []
    resumo = []

    # Capacitações 1 a 4 (legado Power BI)
    for c in legado["capacitacoes"]:
        polo = localiza_municipio(c["polo"])
        capacitacoes.append({"numero": c["numero"], "polo": polo["nome"], "polo_cod_ibge": polo["cod_ibge"],
                             "uf": polo["uf"], "data_inicio": c["data_inicio"], "data_fim": c["data_fim"],
                             "local": None, "fonte": "legado_power_bi"})
    for r in legado["municipio_capacitacao"]:
        fato.append({"capacitacao": r["capacitacao"], "cod_ibge": r["cod_ibge"], "convidado": r["convidado"],
                     "inscritos_total": None, "inscritos_gestores": None, "inscritos_cacs": None, "inscritos_outros": None,
                     "presentes_total": r["presentes_total"], "presentes_gestores": r["presentes_gestores"],
                     "presentes_cacs": r["presentes_cacs"], "presentes_outros": r["presentes_outros"]})

    # Capacitações 5 a 16 (arquivos brutos)
    carrega_referencia_global()
    sem_municipio = {}
    for n in range(5, 17):
        conv = le_convocacao(n)
        convidados = {}
        for txt in conv["convidados_texto"]:
            m = localiza_municipio(txt, UFS_CAPACITACAO[n])
            if not m:
                aviso(f"{n}CTE [Convocação]: município convidado não identificado: '{txt}'")
                continue
            convidados[m["cod_ibge"]] = m
        agg, n_insc, n_pres, n_sem = consolida_capacitacao(n, set(convidados))
        sem_municipio[n] = n_sem
        polo_nome = {5: "Aparecida de Goiânia", 6: "Dourados", 7: "Jardim", 8: "Sinop", 9: "Juína", 10: "Nova Xavantina",
                     11: "Uruaçu", 12: "Brasília", 13: "Alta Floresta", 14: "Jataí", 15: "Posse", 16: "Pontes e Lacerda"}[n]
        polo = localiza_municipio(polo_nome, UFS_CAPACITACAO[n])
        capacitacoes.append({"numero": n, "polo": polo["nome"], "polo_cod_ibge": polo["cod_ibge"], "uf": polo["uf"],
                             "data_inicio": conv["data_inicio"], "data_fim": conv["data_fim"], "local": conv["local"],
                             "fonte": "arquivos_brutos"})
        for cod in sorted(set(convidados) | set(agg)):
            a = agg.get(cod, {"inscritos": Counter(), "presentes": Counter()})
            fato.append({"capacitacao": n, "cod_ibge": cod, "convidado": int(cod in convidados),
                         "inscritos_total": sum(a["inscritos"].values()), "inscritos_gestores": a["inscritos"]["Gestores"],
                         "inscritos_cacs": a["inscritos"]["CACS"], "inscritos_outros": a["inscritos"]["Outros"],
                         "presentes_total": sum(a["presentes"].values()), "presentes_gestores": a["presentes"]["Gestores"],
                         "presentes_cacs": a["presentes"]["CACS"], "presentes_outros": a["presentes"]["Outros"]})
        pres_mun = sum(sum(a["presentes"].values()) for a in agg.values())
        resumo.append((n, len(conv["convidados_texto"]), len(convidados), n_insc, n_pres, pres_mun,
                       sum(1 for c in convidados if sum(agg.get(c, {"presentes": Counter()})["presentes"].values()) > 0),
                       sum(1 for c in agg if c not in convidados and sum(agg[c]["presentes"].values()) > 0)))

    capacitacoes.sort(key=lambda c: c["numero"])
    polo_por_cap = {c["numero"]: c["polo_cod_ibge"] for c in capacitacoes}
    for r in fato:
        r["distancia_polo_km"] = distancia_km(cent.get(r["cod_ibge"]), cent.get(polo_por_cap[r["capacitacao"]]))
    fato.sort(key=lambda r: (r["capacitacao"], r["cod_ibge"]))

    # Avaliações
    av15 = legado["avaliacao_1a5"]["linhas"] + avaliacao_5()
    perguntas, idx_perg = [], {}
    respondentes, itens = {}, []
    for idr, cap, pub, bloco, sub, perg, resp in av15:
        k = (bloco, sub, perg)
        if k not in idx_perg:
            idx_perg[k] = len(perguntas) + 1
            perguntas.append([idx_perg[k], bloco, sub, perg])
        respondentes.setdefault(idr, [idr, cap, pub])
        itens.append([idr, idx_perg[k], resp])
    resp616, notas616, canais616 = avaliacao_6a16()

    cods_usados = {r["cod_ibge"] for r in fato} | {c["polo_cod_ibge"] for c in capacitacoes} | {r[2] for r in resp616 if r[2]}
    municipios = [[c, MUNICIPIOS[c]["nome"], MUNICIPIOS[c]["uf"], *(cent.get(c) or (None, None))] for c in sorted(cods_usados)]

    cols_fato = ["capacitacao", "cod_ibge", "convidado", "inscritos_total", "inscritos_gestores", "inscritos_cacs",
                 "inscritos_outros", "presentes_total", "presentes_gestores", "presentes_cacs", "presentes_outros", "distancia_polo_km"]
    cols_cap = ["numero", "polo", "polo_cod_ibge", "uf", "data_inicio", "data_fim", "local", "fonte"]
    dados = {
        "versao_esquema": 1,
        "gerado_em": datetime.now().strftime("%Y-%m-%d %H:%M"),
        "fontes": {
            "capacitacoes_5_a_16": "data/Capacitação/XCTE - (Convocação | Formulario | Lista de presença | Avaliação)",
            "capacitacoes_1_a_4": "data/Capacitação/legado_cte_01a04_powerbi.json (modelo do Power BI 'bi dados.pbix')",
            "municipios": "IBGE — Lista de Municípios; centróides calculados da malha municipal",
        },
        "presentes_sem_municipio": sem_municipio,
        "capacitacoes": tabela(cols_cap, [[c[k] for k in cols_cap] for c in capacitacoes]),
        "municipios": tabela(["cod_ibge", "nome", "uf", "lat", "lon"], municipios),
        "municipio_capacitacao": tabela(cols_fato, [[r[k] for k in cols_fato] for r in fato]),
        "aval15_perguntas": tabela(["id_pergunta", "bloco", "subgrupo", "pergunta"], perguntas),
        "aval15_respondentes": tabela(["id_resposta", "capacitacao", "publico"], list(respondentes.values())),
        "aval15_respostas": tabela(["id_resposta", "id_pergunta", "resposta"], itens),
        "aval616_perguntas": tabela(["ordem", "pergunta"], [[i + 1, p] for i, p in enumerate(PERGUNTAS_6A16)]),
        "aval616_respondentes": tabela(["id_resposta", "capacitacao", "cod_ibge", "municipio", "cargo", "instituicao",
                                        "aspectos_positivos", "aspectos_melhorar", "temas_futuros", "comentarios"], resp616),
        "aval616_notas": tabela(["id_resposta", "ordem", "nota"], notas616),
        "aval616_canais": tabela(["id_resposta", "canal"], canais616),
    }

    gravar_banco(dados)
    print("CTE  convid(html) convid(ok) inscritos presentes pres_c/mun mun_part mun_espont")
    for r in resumo:
        print("%3d %12d %10d %9d %9d %10d %8d %10d" % r)
    print(f"Avaliação 1-5: {len(respondentes)} respondentes, {len(itens)} respostas; "
          f"6-16: {len(resp616)} respondentes, {len(notas616)} notas, {len(canais616)} canais")
    if avisos:
        print(f"\n{len(avisos)} aviso(s):")
        for a in avisos:
            print(" -", a)


if __name__ == "__main__":
    sys.exit(main())
